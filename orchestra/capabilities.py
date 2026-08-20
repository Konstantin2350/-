import io
import json
import math
import re
import struct
import wave
from collections import defaultdict
from datetime import UTC, date, datetime
from typing import Any

import httpx
from pptx import Presentation

from orchestra.config import Settings
from orchestra.schemas import (
    ContentRequest,
    KPIForecastRequest,
    ProcessEvent,
    ProcessNLRequest,
    ProjectDigestRequest,
    TrainingEvaluateRequest,
)
from orchestra.services import Embeddings, LLMClient, tokens


class ExternalProviderRequired(RuntimeError):
    pass


class AudioService:
    audio_extensions = {".mp3", ".wav", ".m4a", ".ogg", ".webm", ".flac", ".mp4"}

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    async def transcribe(self, filename: str, content: bytes, language: str | None = None) -> str:
        extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if extension not in self.audio_extensions:
            raise ValueError("Supported audio formats: MP3, WAV, M4A, OGG, WEBM, FLAC, MP4")
        if len(content) > self.settings.max_audio_bytes:
            raise ValueError("Audio exceeds configured upload limit")
        key = self.settings.stt_api_key or self.settings.llm_api_key
        if not key:
            raise ExternalProviderRequired(
                "Configure STT_API_KEY or LLM_API_KEY for audio transcription"
            )
        base_url = self.settings.stt_base_url or self.settings.llm_base_url
        data = {"model": self.settings.stt_model}
        if language:
            data["language"] = language
        async with httpx.AsyncClient(timeout=300) as client:
            response = await client.post(
                f"{base_url.rstrip('/')}/audio/transcriptions",
                headers={"Authorization": f"Bearer {key}"},
                data=data,
                files={"file": (filename, content, "application/octet-stream")},
            )
            response.raise_for_status()
            payload = response.json()
        return payload["text"]

    async def synthesize(self, text: str, voice: str, output_format: str) -> bytes:
        if not self.settings.llm_api_key:
            raise ExternalProviderRequired("Configure LLM_API_KEY for speech synthesis")
        async with httpx.AsyncClient(timeout=180) as client:
            response = await client.post(
                f"{self.settings.llm_base_url.rstrip('/')}/audio/speech",
                headers={"Authorization": f"Bearer {self.settings.llm_api_key}"},
                json={
                    "model": "gpt-4o-mini-tts",
                    "voice": voice,
                    "input": text,
                    "response_format": output_format,
                },
            )
            response.raise_for_status()
            return response.content

    @staticmethod
    def wav_signals(content: bytes, transcript: str = "") -> dict[str, Any]:
        try:
            with wave.open(io.BytesIO(content), "rb") as audio:
                frames = audio.getnframes()
                rate = audio.getframerate()
                width = audio.getsampwidth()
                channels = audio.getnchannels()
                raw = audio.readframes(min(frames, rate * 10))
        except (wave.Error, EOFError):
            return {"available": False, "reason": "Prosody analysis requires PCM WAV"}
        if not frames or not rate or width not in {1, 2, 4}:
            return {"available": False, "reason": "Unsupported WAV encoding"}
        format_code = {1: "B", 2: "h", 4: "i"}[width]
        sample_count = len(raw) // width
        values = struct.unpack(f"<{sample_count}{format_code}", raw)
        if width == 1:
            values = tuple(value - 128 for value in values)
        mono = values[::channels]
        max_value = float(2 ** (8 * width - 1))
        rms = math.sqrt(sum(value * value for value in mono) / max(len(mono), 1))
        silence = sum(abs(value) < max_value * 0.02 for value in mono)
        duration = frames / rate
        words_per_minute = (
            round(len(tokens(transcript)) / duration * 60, 1) if transcript and duration else None
        )
        return {
            "available": True,
            "duration_seconds": round(duration, 2),
            "energy": round(rms / max_value, 4),
            "silence_ratio": round(silence / max(len(mono), 1), 3),
            "words_per_minute": words_per_minute,
            "pace": (
                "fast"
                if words_per_minute and words_per_minute > 170
                else "slow"
                if words_per_minute and words_per_minute < 90
                else "normal"
                if words_per_minute
                else "unknown"
            ),
        }


class ProcessService:
    allowed_step_types = {"agent", "bitrix_call", "condition", "wait", "notify", "human_approval"}

    def compile_nl(self, request: ProcessNLRequest) -> dict[str, Any]:
        description = request.description.strip()
        trigger_match = re.search(r"(?:когда|если|при)\s+(.+?)(?:,|;|\n|то\s)", description, re.I)
        trigger = trigger_match.group(1).strip() if trigger_match else "manual"
        body = description[trigger_match.end() :] if trigger_match else description
        phrases = [
            phrase.strip(" .,-")
            for phrase in re.split(r"\b(?:затем|после этого|потом)\b|[;\n]", body, flags=re.I)
            if phrase.strip(" .,-")
        ]
        steps = [self._step(phrase, index) for index, phrase in enumerate(phrases)]
        return {
            "dsl_version": "1.0",
            "process": {
                "name": description[:80],
                "trigger": trigger,
                "steps": steps,
            },
            "validation": {"valid": bool(steps), "steps": len(steps)},
        }

    @staticmethod
    def _step(phrase: str, index: int) -> dict[str, Any]:
        lowered = phrase.lower()
        if any(word in lowered for word in ("согласовать", "утвердить", "подтвердить")):
            step_type = "human_approval"
        elif any(word in lowered for word in ("уведом", "сообщ", "отправить письмо")):
            step_type = "notify"
        elif any(word in lowered for word in ("подожд", "через ", "спустя")):
            step_type = "wait"
        elif any(word in lowered for word in ("bitrix", "сделк", "лид", "crm")):
            step_type = "bitrix_call"
        elif lowered.startswith(("если ", "иначе ")):
            step_type = "condition"
        else:
            step_type = "agent"
        return {"id": f"step_{index + 1}", "type": step_type, "instruction": phrase}

    def validate(self, steps: list[dict[str, Any]]) -> list[int]:
        return [
            index
            for index, step in enumerate(steps)
            if step.get("type") not in self.allowed_step_types
        ]


class ContentService:
    def __init__(self, llm: LLMClient) -> None:
        self.llm = llm

    async def generate(self, request: ContentRequest) -> dict[str, Any]:
        generated = await self.llm.complete(
            (
                f"Создай {request.kind}. Аудитория: {request.audience}. "
                f"Тон: {request.tone}. Не выдумывай факты."
            ),
            request.brief,
            temperature=0.5,
        )
        if generated:
            return {"content": generated, "provider": "llm"}
        return {
            "content": self._fallback(request),
            "provider": "local-template",
        }

    @staticmethod
    def _fallback(request: ContentRequest) -> str:
        if request.kind == "meta":
            title = request.brief.split(".", 1)[0][:60]
            return json.dumps(
                {
                    "title": title,
                    "description": request.brief[:160],
                    "keywords": list(dict.fromkeys(tokens(request.brief)))[:8],
                },
                ensure_ascii=False,
            )
        if request.kind == "brainstorm":
            subject = request.brief.strip()
            return "\n".join(
                [
                    f"1. Упростить: как получить результат «{subject}» за один шаг?",
                    f"2. Автоматизировать: какие повторяемые части «{subject}» убрать?",
                    f"3. Измерить: какой показатель подтвердит успех «{subject}»?",
                    f"4. Масштабировать: что позволит повторить «{subject}» для 100 клиентов?",
                ]
            )
        if request.kind == "meeting_summary":
            sentences = [
                part.strip() for part in re.split(r"(?<=[.!?])\s+", request.brief) if part.strip()
            ]
            actions = [
                sentence
                for sentence in sentences
                if re.search(r"\b(нужно|отправ|сдела|подготов|до |завтра)\w*", sentence, re.I)
            ]
            return (
                "Резюме: "
                + " ".join(sentences[:3])
                + "\nДействия:\n"
                + "\n".join(f"- {item}" for item in actions[:10])
            )
        return (
            f"Тема: {request.brief}\n"
            f"Аудитория: {request.audience}\n\n"
            "Проблема: опишите текущую ситуацию клиента.\n"
            "Решение: покажите конкретную пользу и подтверждённые факты.\n"
            "Следующий шаг: предложите одно понятное действие."
        )

    @staticmethod
    def presentation(title: str, content: str, slide_count: int) -> bytes:
        deck = Presentation()
        first = deck.slides.add_slide(deck.slide_layouts[0])
        first.shapes.title.text = title
        first.placeholders[1].text = "Подготовлено ИИ‑Оркестром"
        points = [
            part.strip(" \n-")
            for part in re.split(r"[\n]+|(?<=[.!?])\s+", content)
            if len(part.strip()) > 10
        ]
        per_slide = max(1, math.ceil(len(points) / max(slide_count - 1, 1)))
        for index in range(0, len(points), per_slide):
            if len(deck.slides) >= slide_count:
                break
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = f"{title}: часть {len(deck.slides)}"
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for point_index, point in enumerate(points[index : index + per_slide]):
                paragraph = frame.paragraphs[0] if point_index == 0 else frame.add_paragraph()
                paragraph.text = point[:500]
        output = io.BytesIO()
        deck.save(output)
        return output.getvalue()


class TrainingService:
    def __init__(self, embeddings: Embeddings) -> None:
        self.embeddings = embeddings

    @staticmethod
    def generate(source: str, count: int) -> list[dict[str, Any]]:
        sentences = [
            part.strip()
            for part in re.split(r"(?<=[.!?])\s+|\n+", source)
            if len(tokens(part)) >= 6
        ]
        questions = []
        for index, sentence in enumerate(sentences[:count]):
            subject = " ".join(tokens(sentence)[:6])
            questions.append(
                {
                    "id": index + 1,
                    "question": f"Что в материале сказано про «{subject}»?",
                    "expected_answer": sentence,
                    "difficulty": "medium" if len(tokens(sentence)) > 15 else "easy",
                }
            )
        return questions

    async def evaluate(self, request: TrainingEvaluateRequest) -> dict[str, Any]:
        vectors = await self.embeddings.encode_many([request.expected, request.answer])
        semantic = max(0.0, Embeddings.similarity(vectors[0], vectors[1]))
        expected_words = set(tokens(request.expected))
        answer_words = set(tokens(request.answer))
        lexical = len(expected_words & answer_words) / max(len(expected_words), 1)
        score = round(min(1.0, 0.55 * semantic + 0.45 * lexical), 3)
        recent = (request.previous_scores + [score])[-5:]
        average = sum(recent) / len(recent)
        next_difficulty = "hard" if average >= 0.8 else "easy" if average < 0.5 else "medium"
        return {
            "score": score,
            "passed": score >= 0.65,
            "next_difficulty": next_difficulty,
            "feedback": (
                "Ответ принят."
                if score >= 0.65
                else "Повторите материал и добавьте ключевые факты из эталона."
            ),
        }


class AnalyticsService:
    @staticmethod
    def process_mining(events: list[ProcessEvent]) -> dict[str, Any]:
        cases: dict[str, list[ProcessEvent]] = defaultdict(list)
        for event in events:
            cases[event.case_id].append(event)
        durations: dict[str, list[float]] = defaultdict(list)
        variants: dict[tuple[str, ...], int] = defaultdict(int)
        for case_events in cases.values():
            ordered = sorted(case_events, key=lambda item: item.occurred_at)
            variants[tuple(item.activity for item in ordered)] += 1
            for current, following in zip(ordered, ordered[1:], strict=False):
                elapsed = (following.occurred_at - current.occurred_at).total_seconds()
                durations[current.activity].append(max(0, elapsed))
        averages = {
            activity: round(sum(values) / len(values), 2) for activity, values in durations.items()
        }
        bottlenecks = sorted(
            (
                {"activity": activity, "average_seconds": duration}
                for activity, duration in averages.items()
            ),
            key=lambda item: item["average_seconds"],
            reverse=True,
        )
        return {
            "cases": len(cases),
            "variants": [
                {"path": list(path), "cases": count}
                for path, count in sorted(variants.items(), key=lambda item: item[1], reverse=True)[
                    :20
                ]
            ],
            "bottlenecks": bottlenecks[:10],
            "recommendations": [
                f"Проверьте этап «{item['activity']}»: среднее ожидание "
                f"{item['average_seconds'] / 3600:.1f} ч."
                for item in bottlenecks[:3]
            ],
        }

    @staticmethod
    def forecast(request: KPIForecastRequest) -> dict[str, Any]:
        values = request.values
        count = len(values)
        x_mean = (count - 1) / 2
        y_mean = sum(values) / count
        denominator = sum((index - x_mean) ** 2 for index in range(count)) or 1
        slope = (
            sum((index - x_mean) * (value - y_mean) for index, value in enumerate(values))
            / denominator
        )
        intercept = y_mean - slope * x_mean
        fitted = [intercept + slope * index for index in range(count)]
        mae = (
            sum(abs(actual - predicted) for actual, predicted in zip(values, fitted, strict=True))
            / count
        )
        forecast = [
            round(intercept + slope * index, 3) for index in range(count, count + request.horizon)
        ]
        return {
            "forecast": forecast,
            "trend_per_period": round(slope, 4),
            "mae": round(mae, 4),
            "method": "linear_trend",
        }

    @staticmethod
    def project_digest(request: ProjectDigestRequest) -> dict[str, Any]:
        today = datetime.now(UTC).date()
        overdue = [
            task
            for task in request.tasks
            if task.status != "done" and task.due_date and task.due_date < today
        ]
        blocked = [task for task in request.tasks if task.status == "blocked"]
        at_risk = [
            task
            for task in request.tasks
            if task.status != "done"
            and task.due_date
            and task.due_date <= today
            and task.progress < 80
        ]
        completion = (
            sum(task.progress for task in request.tasks) / len(request.tasks)
            if request.tasks
            else 0
        )
        return {
            "project": request.project,
            "generated_at": date.today().isoformat(),
            "completion_percent": round(completion, 1),
            "overdue": [task.model_dump(mode="json") for task in overdue],
            "blocked": [task.model_dump(mode="json") for task in blocked],
            "at_risk": [task.model_dump(mode="json") for task in at_risk],
            "summary": (
                f"Готовность {completion:.0f}%. Просрочено: {len(overdue)}, "
                f"заблокировано: {len(blocked)}, под риском: {len(at_risk)}."
            ),
        }
